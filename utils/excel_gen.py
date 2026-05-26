import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font

def generate_report(users, file_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Loyiha Hisoboti"

    # Header
    columns = [
        "No", "Invoist ID", "Mulk egasi", "Manzili", "Tel rakami", 
        "Loyiha toifasi", "Xizmat turi", "Bosqich", "Turi", 
        "Qiymati", "Tulangan", "Qoldiq", "To'lov Usuli", "Ish xolat", "Vaqt"
    ]
    
    ws.append(columns)

    # Styles
    thin_border = Border(
        left=Side(style='thin', color='808080'),
        right=Side(style='thin', color='808080'),
        top=Side(style='thin', color='808080'),
        bottom=Side(style='thin', color='808080')
    )
    center_aligned = Alignment(horizontal='center', vertical='center')
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Paid
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")     # Remaining
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Status

    # Style Header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = center_aligned
        cell.border = thin_border

    # Data
    for i, user in enumerate(users, start=1):
        total = user['total_sum'] or 0
        paid = user['paid_sum'] or 0
        balance = total - paid
        
        row = [
            i, 
            user['tg_id'], 
            user['full_name'], 
            user['address'], 
            user['phone'],
            user['loyiha_toifasi'], 
            user['service_type'], 
            user['stage'], 
            user['tur'],
            total, 
            paid, 
            balance, 
            user.get('payment_method', 'Noma\'lum'),
            user['status'], 
            user['created_at']
        ]
        ws.append(row)
        
        # Style row cells
        current_row = ws.max_row
        for col_idx, cell in enumerate(ws[current_row], start=1):
            cell.alignment = center_aligned
            cell.border = thin_border
            
            # Apply column-specific formatting
            if col_idx in [10, 11, 12]:
                cell.number_format = '#,##0'
            
            # Color fills
            if col_idx == 11: # Tulangan
                cell.fill = green_fill
            elif col_idx == 12: # Qoldiq
                cell.fill = red_fill
            elif col_idx == 14: # Ish xolat
                cell.fill = yellow_fill

    # Adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = max_length + 2

    wb.save(file_path)
